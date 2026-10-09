import json
import os
import streamlit as st
import ollama
from pptx import Presentation
from dotenv import load_dotenv
from powerpointEngine import PowerPointMCPEngine
load_dotenv()

# class PowerPointMCPEngine:
#     def __init__(self):
#         self.prs = Presentation()
#
#     def add_title_slide(self, title : str, subtitle: str = ""):
#         slide_layout = self.prs.slide_layouts[0]
#         slide = self.prs.slides.add_slide(slide_layout)
#         slide.shapes.title.text = title
#         if subtitle and len(slide.placeholders) > 1:
#             slide.placeholders[1].text = subtitle
#         return f"Added Title Slide: '{title}'"
#
#     def add_content_slide(self, title: str, bullet_points : list):
#         slide_layout = self.prs.slide_layouts[1]
#         slide = self.prs.slides.add_slide(slide_layout)
#         slide.shapes.title.text = title
#         tf = slide.placeholders[1].text_frame
#         for i, point in enumerate(bullet_points):
#             p = tf.add_paragraph() if i > 0 else tf.paragraphs[0]
#             p.text = point
#         return f"Added Content Slide: {title} with {len(bullet_points)} items"
#
#     def save_presentation(self, filename : str):
#         if not filename.endswith('.pptx'):
#             filename += ".pptx"
#         output_path = os.path.join(os.getcwd(), filename)
#         self.prs.save(output_path)
#         return output_path
#
#     def execute_tool(self, tool_name: str, args: dict):
#         if tool_name == "add_title_slide":
#             return self.add_title_slide(args.get("title", ""), args.get("subtitle", ""))
#         elif tool_name == "add_content_slide":
#             return self.add_content_slide(args.get('title', ''), args.get('bullet_points', []))
#         elif tool_name == "save_presentation":
#             return self.save_presentation(args.get("filename", "presentation.pptx"))
#         else:
#             raise ValueError(f"Unknown tool name: {tool_name}")
#

st.set_page_config(page_title= "AI Powerpoint Generator", page_icon="📊", layout = 'wide')
st.title("📊AI Powerpoint Generator")
st.caption(f"Powered by local **Ollama {os.getenv("MODEL_NAME")}** & **Powerpoint MCP Engine**")

with st.sidebar:
    st.header("⚙️ Configuration")
    model_name = st.text_input("Ollama Model", value = os.getenv("MODEL_NAME"))
    filename_input = st.text_input("Output Filename", value = "generated_presentation.pptx")

    st.markdown("---")
    st.markdown("### Available MCP Tools")
    st.code("""
    - add_title_slide(title, subtitle)
    - add_content_slide(title, bullet_points)
    - save_presentation(filename)
        """, language="text")

user_prompt = st.text_area(
        "Describe the presentation you want to build:",
        height=120,
        placeholder="Create a 4-slide presentation about Machine Learning basics, application in healthcare, and future trends."
    )

generate_btn = st.button("🚀 Generate Presentation", type="primary", use_container_width=True)


# SYSTEM_PROMPT = """You are an expert executive presentation creator. Your task is to design detailed, content-rich PowerPoint presentation outlines using available tool calls.
#
# Available Tools Schema:
# 1. {"tool": "add_title_slide", "args": {"title": "Slide Title", "subtitle": "Detailed Subtitle explaining the presentation goals"}}
# 2. {"tool": "add_content_slide", "args": {"title": "Specific Slide Title", "bullet_points": ["Detailed point 1", "Detailed point 2"]}}
# 3. {"tool": "save_presentation", "args": {"filename": "presentation.pptx"}}
#
# Rules for Content Quality:
# - EVERY bullet point must be a FULL, INFORMATIVE sentence (15–25 words per point). Do NOT use 2-3 word phrases like "Document Search".
# - Include specific explanations, metrics, key benefits, and real-world context inside each bullet point.
# - Provide 3 to 4 comprehensive bullet points per content slide.
# - Structure slides logically: Start with 'add_title_slide', add multiple detailed 'add_content_slide' calls, and end with 'save_presentation'.
# - Return strictly a valid JSON array of tool calls matching the schema above.
# """

# Extended System Prompt enforcing visual layout tools
SYSTEM_PROMPT = """You are an expert presentation designer. Convert user prompts into structured JSON tool calls.

Available Tools Schema:
1. {"tool": "add_title_slide", "args": {"title": "Main Title", "subtitle": "Supporting Subtitle"}}
2. {"tool": "add_content_slide", "args": {"title": "Slide Title", "bullet_points": ["Detailed Point 1", "Detailed Point 2", "Detailed Point 3"]}}
3. {"tool": "add_stat_slide", "args": {"title": "Slide Title", "stat_number": "85%", "description": "Full explanation of the metric impact."}}
4. {"tool": "save_presentation", "args": {"filename": "presentation.pptx"}}

Rules:
- Provide 2 to 4 detailed bullet points for 'add_content_slide'.
- Use 'add_stat_slide' whenever numerical data, KPIs, or percentages are mentioned.
- Always begin with 'add_title_slide' and end with 'save_presentation'.
"""

schema = {
    "type": "array",
    "items": {
        "type": "object",
        "properties": {
            "tool": {"type": "string"},
            "args": {"type": "object"}
        },
        "required": ["tool", "args"]
    }
}

if generate_btn:
    if not user_prompt.strip():
        st.warning("Please enter a prompt first!")
    else:
        engine = PowerPointMCPEngine()
        status_container = st.status("Initializing generation process...", expanded = True)

        try:
            status_container.write(f"🤖 **Step 1:** Querying local Ollama model (`{os.getenv("MODEL_NAME")}`)...")

            response = ollama.generate(
                model = model_name,
                prompt = f"{SYSTEM_PROMPT}\nUser Prompt: {user_prompt}\nJSON Output:",
                format= schema
            )

            raw_response = response.get("response", "")
            actions = json.loads(raw_response)

            status_container.write("🧠 **Step 2:** Plan generated from Ollama!")

            st.subheader("Execution Plan")
            st.json(actions)

            status_container.write("⚙️ **Step 3:** Executing MCP commands to build PPTX...")
            logs = []
            saved_file_path = None

            for action in actions:
                tool_name = action.get("tool")
                args = action.get("args", {})

                if tool_name == "save_presentation" and filename_input:
                    args["filename"] = filename_input

                res = engine.execute_tool(tool_name, args)
                logs.append(res)


            status_container.update(label="✅ Presentation generated successfully!", state="complete", expanded=False)
            st.success("🎉 PowerPoint deck ready for download!")

            if saved_file_path and os.path.exists(saved_file_path):
                with open(saved_file_path, "rb") as file:
                    st.download_button(
                        label="📥 Download PowerPoint (.pptx)",
                        data=file,
                        file_name=os.path.basename(saved_file_path),
                        mime="application/vnd.openxmlformats-officedocument.presentationml.presentation",
                        type = 'primary'
                    )


        except json.JSONDecodeError as e:
            status_container.update(label = "❌ Failed to parse JSON from Ollama output", state = 'error')
            st.error(f"Ollama did not return structured JSON. Respone was: \n```\n{raw_response}\n```")

        except Exception as e:
            status_container.update(label = "❌ An error occured during generation", state="error")
            st.error(f"Error details: {str(e)}")




















