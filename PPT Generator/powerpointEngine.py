from pptx import Presentation
import os
from pptx.util import Inches, Pt
from pptx.enum.shapes import MSO_SHAPE
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN


# class PowerPointMCPEngine:
#     def __init__(self):
#         self.prs = Presentation()
#
#     def add_title_slide(self, title: str, subtitle: str = ""):
#         slide_layout = self.prs.slide_layouts[0]
#         slide = self.prs.slides.add_slide(slide_layout)
#         slide.shapes.title.text = title
#         if subtitle and len(slide.placeholders) > 1:
#             slide.placeholders[1].text = subtitle
#         return f"Added Title Slide: '{title}'"
#
#     def add_content_slide(self, title: str, bullet_points: list):
#         slide_layout = self.prs.slide_layouts[1]
#         slide = self.prs.slides.add_slide(slide_layout)
#         slide.shapes.title.text = title
#         tf = slide.placeholders[1].text_frame
#         for i, point in enumerate(bullet_points):
#             p = tf.add_paragraph() if i > 0 else tf.paragraphs[0]
#             p.text = point
#         return f"Added Content Slide: {title} with {len(bullet_points)} items"
#
#     def save_presentation(self, filename: str):
#         if not filename.endswith('.pptx'):
#             filename += ".pptx"
#         output_path = os.path.join(os.getcwd() + "\\presentations", filename)
#         self.prs.save(output_path)
#         return output_path
#
#     def execute_tool(self, tool_name: str, args: dict):
#         if tool_name == "add_title_slide":
#             return self.add_title_slide(args.get("title", ""), args.get("subtitle", ""))
#         elif tool_name == "add_content_slide":
#             return self.add_content_slide(args.get('title', ''), args.get('bullet_points', []))
#         elif tool_name == "save_presentation":
#             return self.save_presentation(args.get("filename", "presentation.pptx"))
#         else:
#             raise ValueError(f"Unknown tool name: {tool_name}")
#


COLOR_PRIMARY = RGBColor(15, 23, 42)      # Deep Navy / Slate-900
COLOR_ACCENT = RGBColor(13, 148, 136)     # Teal-600
COLOR_CARD_BG = RGBColor(241, 245, 249)   # Light Slate-100
COLOR_CARD_BORDER = RGBColor(226, 232, 240) # Border Slate-200
COLOR_TEXT_BODY = RGBColor(51, 65, 85)    # Slate-700
COLOR_MUTED = RGBColor(100, 116, 139)     # Slate-500
COLOR_WHITE = RGBColor(255, 255, 255)


class PowerPointMCPEngine:
    def __init__(self):
        self.prs = Presentation()
        # Set 16:9 Widescreen dimensions
        self.prs.slide_width = Inches(13.333)
        self.prs.slide_height = Inches(7.5)
        # Blank slide layout
        self.blank_layout = self.prs.slide_layouts[6]

    def _add_header(self, slide, title_text: str, category_text: str = "PRESENTATION"):
        """Adds a structured visual header with category tag and accent line."""
        # Category Eyebrow
        cat_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.7), Inches(0.4))
        tf_cat = cat_box.text_frame
        tf_cat.word_wrap = True
        p_cat = tf_cat.paragraphs[0]
        p_cat.text = category_text.upper()
        p_cat.font.size = Pt(11)
        p_cat.font.bold = True
        p_cat.font.color.rgb = COLOR_ACCENT

        # Title
        title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.7), Inches(11.7), Inches(0.8))
        tf_title = title_box.text_frame
        tf_title.word_wrap = True
        p_title = tf_title.paragraphs[0]
        p_title.text = title_text
        p_title.font.size = Pt(26)
        p_title.font.bold = True
        p_title.font.color.rgb = COLOR_PRIMARY

        # Accent Line divider
        accent_line = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(1.5), Inches(1.2), Inches(0.05)
        )
        accent_line.fill.solid()
        accent_line.fill.fore_color.rgb = COLOR_ACCENT
        accent_line.line.color.rgb = COLOR_ACCENT

    def add_title_slide(self, title: str, subtitle: str = ""):
        """Creates a modern hero title slide with dark background."""
        slide = self.prs.slides.add_slide(self.blank_layout)

        # Full dark background shape
        bg = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5)
        )
        bg.fill.solid()
        bg.fill.fore_color.rgb = COLOR_PRIMARY
        bg.line.fill.background()

        # Accent Bar
        accent = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, Inches(1.2), Inches(2.2), Inches(0.8), Inches(0.08)
        )
        accent.fill.solid()
        accent.fill.fore_color.rgb = COLOR_ACCENT
        accent.line.fill.background()

        # Title Text Box
        title_box = slide.shapes.add_textbox(Inches(1.2), Inches(2.5), Inches(10.5), Inches(2.0))
        tf = title_box.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = title
        p.font.size = Pt(44)
        p.font.bold = True
        p.font.color.rgb = COLOR_WHITE

        # Subtitle Text Box
        if subtitle:
            sub_box = slide.shapes.add_textbox(Inches(1.2), Inches(4.5), Inches(10.5), Inches(1.5))
            tf_sub = sub_box.text_frame
            tf_sub.word_wrap = True
            p_sub = tf_sub.paragraphs[0]
            p_sub.text = subtitle
            p_sub.font.size = Pt(20)
            p_sub.font.color.rgb = COLOR_CARD_BG

        return f"Added Visual Title Slide: '{title}'"

    def add_content_slide(self, title: str, bullet_points: list):
        """Creates a side-by-side modular card layout based on point count."""
        slide = self.prs.slides.add_slide(self.blank_layout)
        self._add_header(slide, title, category_text="KEY INSIGHTS")

        num_cards = max(1, min(len(bullet_points), 4))  # Cap at 4 cards max
        total_width = 11.733  # Total width for cards container
        gap = 0.35
        card_width = (total_width - (gap * (num_cards - 1))) / num_cards
        card_top = Inches(1.8)
        card_height = Inches(4.8)

        for i in range(num_cards):
            left_pos = Inches(0.8 + i * (card_width + gap))

            # Draw Card Background Box
            card_shape = slide.shapes.add_shape(
                MSO_SHAPE.ROUNDED_RECTANGLE, left_pos, card_top, Inches(card_width), card_height
            )
            card_shape.fill.solid()
            card_shape.fill.fore_color.rgb = COLOR_CARD_BG
            card_shape.line.color.rgb = COLOR_CARD_BORDER
            card_shape.line.width = Pt(1)

            # Add Card Header/Badge
            badge = slide.shapes.add_shape(
                MSO_SHAPE.OVAL, left_pos + Inches(0.25), card_top + Inches(0.3), Inches(0.4), Inches(0.4)
            )
            badge.fill.solid()
            badge.fill.fore_color.rgb = COLOR_ACCENT
            badge.line.fill.background()

            # Badge Number Text
            tf_badge = badge.text_frame
            p_b = tf_badge.paragraphs[0]
            p_b.text = str(i + 1)
            p_b.alignment = PP_ALIGN.CENTER
            p_b.font.size = Pt(14)
            p_b.font.bold = True
            p_b.font.color.rgb = COLOR_WHITE

            # Card Text Box
            text_box = slide.shapes.add_textbox(
                left_pos + Inches(0.25), card_top + Inches(0.9), Inches(card_width - 0.5), Inches(3.6)
            )
            tf = text_box.text_frame
            tf.word_wrap = True

            point_text = bullet_points[i]
            p = tf.paragraphs[0]
            p.text = point_text
            p.font.size = Pt(15)
            p.font.color.rgb = COLOR_TEXT_BODY
            p.line_spacing = 1.2

        return f"Added Card Layout Slide: '{title}' with {num_cards} cards"

    def add_stat_slide(self, title: str, stat_number: str, description: str):
        """Creates a high-impact metric callout slide."""
        slide = self.prs.slides.add_slide(self.blank_layout)
        self._add_header(slide, title, category_text="METRIC CALLOUT")

        # Big Stat Card Box
        card = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.8), Inches(11.733), Inches(4.8)
        )
        card.fill.solid()
        card.fill.fore_color.rgb = COLOR_CARD_BG
        card.line.color.rgb = COLOR_CARD_BORDER

        # Stat Number
        num_box = slide.shapes.add_textbox(Inches(1.2), Inches(2.2), Inches(10.9), Inches(1.8))
        tf_num = num_box.text_frame
        p_num = tf_num.paragraphs[0]
        p_num.text = stat_number
        p_num.font.size = Pt(72)
        p_num.font.bold = True
        p_num.font.color.rgb = COLOR_ACCENT

        # Description Label
        desc_box = slide.shapes.add_textbox(Inches(1.2), Inches(4.2), Inches(10.9), Inches(1.8))
        tf_desc = desc_box.text_frame
        tf_desc.word_wrap = True
        p_desc = tf_desc.paragraphs[0]
        p_desc.text = description
        p_desc.font.size = Pt(20)
        p_desc.font.color.rgb = COLOR_TEXT_BODY

        return f"Added Stat Slide: '{stat_number}'"

    def save_presentation(self, filename: str):
        if not filename.endswith(".pptx"):
            filename += ".pptx"
        output_path = os.path.join(os.getcwd() + "\\presentations", filename)
        self.prs.save(output_path)
        return output_path

    def execute_tool(self, tool_name: str, args: dict):
        if tool_name == "add_title_slide":
            return self.add_title_slide(args.get("title", ""), args.get("subtitle", ""))
        elif tool_name == "add_content_slide":
            return self.add_content_slide(args.get("title", ""), args.get("bullet_points", []))
        elif tool_name == "add_stat_slide":
            return self.add_stat_slide(args.get("title", ""), args.get("stat_number", ""), args.get("description", ""))
        elif tool_name == "save_presentation":
            return self.save_presentation(args.get("filename", "presentation.pptx"))
        else:
            raise ValueError(f"Unknown tool name: {tool_name}")